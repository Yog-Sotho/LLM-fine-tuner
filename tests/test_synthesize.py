"""Training data from documents: reading, chunking, Q/A writing, curation, UI and CLI.

The writer is a real HTTP server speaking the OpenAI chat API (in a thread), so the
remote path runs end to end without a network or a large model.
"""

import json
import re
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest
from typer.testing import CliRunner

import cli.commands as commands
from data.documents import chunk_text, document_chunks, read_document
from inference.synthesize import format_stats, parse_qa_pairs, synthesize_pairs

# ── Documents ──────────────────────────────────────────────────────────────


def test_chunks_cover_the_text_with_overlap_and_clean_cuts():
    text = "\n\n".join(f"Paragraph {i}. " + "Some words here. " * 30 for i in range(10))
    chunks = chunk_text(text, size=800, overlap=100)
    assert len(chunks) > 3 and all(len(c) <= 800 for c in chunks)
    words = set(text.split())
    assert all(c.split()[0] in words and c.split()[-1] in words for c in chunks)  # whole words
    for a, b in zip(chunks, chunks[1:], strict=False):
        assert b[:30] in a  # consecutive chunks overlap
    assert chunk_text("") == [] and chunk_text("short") == ["short"]
    with pytest.raises(ValueError, match="larger than the overlap"):
        chunk_text("x" * 50, size=10, overlap=10)


def test_documents_of_each_type(tmp_path):
    import docx

    (tmp_path / "a.txt").write_text("Plain text file.")
    (tmp_path / "b.md").write_text("# Title\n\nMarkdown body.")
    document = docx.Document()
    document.add_paragraph("Word paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "cell A", "cell B"
    document.save(tmp_path / "c.docx")
    assert read_document(str(tmp_path / "a.txt")) == "Plain text file."
    assert "Markdown body." in read_document(str(tmp_path / "b.md"))
    word = read_document(str(tmp_path / "c.docx"))
    assert "Word paragraph." in word and "cell A | cell B" in word
    chunks = document_chunks([str(tmp_path / n) for n in ("a.txt", "b.md", "c.docx")])
    assert [name for name, _ in chunks] == ["a.txt", "b.md", "c.docx"]


@pytest.mark.parametrize(
    ("name", "message"), [("x.csv", "use one of"), ("missing.pdf", "File not found"),
                          ("../x.txt", "Path traversal")],
)  # fmt: skip
def test_unreadable_documents(tmp_path, name, message):
    if name == "x.csv":
        (tmp_path / name).write_text("a,b")
    path = name if name.startswith("..") else str(tmp_path / name)
    with pytest.raises(ValueError, match=message):
        read_document(path)


# ── Writing and curating pairs ─────────────────────────────────────────────


@pytest.mark.parametrize(
    ("reply", "count"),
    [
        ('[{"question": "Q?", "answer": "A."}]', 1),
        ('Here you go:\n```json\n[{"question": "Q?", "answer": "A."}]\n```', 1),
        ('[{"question": "Q?", "answer": ""}, {"q": 1}, "text"]', 0),
        ("Sorry, I can't.", 0),
        ("[not json", 0),
    ],
)
def test_parse_pairs(reply, count):
    assert len(parse_qa_pairs(reply)) == count


def _fake_writer(rating=8, reply=None, fail_on=()):
    def ask(prompt):
        if prompt.startswith("Rate"):
            return "unclear" if rating is None else f"Score: {rating}"
        chunk = re.search(r'"""\n(.*)\n"""', prompt, re.S).group(1)
        if chunk in fail_on:
            raise RuntimeError("server error hf_" + "c" * 34)
        return reply or json.dumps([{"question": f"What does '{chunk[:12]}' say?", "answer": chunk},
                                    {"question": "Same question?", "answer": "x"}])  # fmt: skip

    return ask


def test_curation_dedup_and_failures():
    chunks = [("a.pdf", "first chunk"), ("a.pdf", "second chunk"), ("b.pdf", "broken")]
    rows, stats = synthesize_pairs(chunks, _fake_writer(8, fail_on={"broken"}), 2, 7)
    questions = [r["instruction"] for r in rows]
    assert len(rows) == 3 and questions.count("Same question?") == 1  # duplicate dropped
    assert stats["failed_chunks"] == 1 and stats["duplicates"] == 1 and rows[0]["score"] == 8
    assert {r["source"] for r in rows} == {"a.pdf"}
    _, low = synthesize_pairs(chunks[:1], _fake_writer(4), 2, 7)
    assert low["kept"] == 0 and low["rejected"] == 2
    _, unrated = synthesize_pairs(chunks[:1], _fake_writer(None), 2, 7)
    assert unrated["kept"] == 0 and unrated["unrated"] == 2
    kept_all, _ = synthesize_pairs(chunks[:1], _fake_writer(1), 2, 0)  # threshold 0: no rating
    assert len(kept_all) == 2 and kept_all[0]["score"] is None


def test_stop_and_progress():
    seen = []
    rows, stats = synthesize_pairs([("a", "one"), ("a", "two")], _fake_writer(), 1, 0,
                                   lambda i, n: seen.append((i, n)), lambda: len(seen) >= 1)  # fmt: skip
    assert seen == [(0, 2)] and stats.get("stopped") and len(rows) == 2


def test_stats_message():
    stats = {"chunks": 2, "generated": 6, "duplicates": 1, "rejected": 2, "failed_chunks": 1,
             "unrated": 0, "kept": 3}  # fmt: skip
    text = format_stats(stats)
    assert text.startswith("✅ 3 question/answer pairs from 2 text chunks")
    assert "2 rated too low" in text and "1 duplicates" in text and "1 chunks gave no" in text
    assert format_stats({**stats, "kept": 0}).startswith("❌")


# ── A real OpenAI-compatible server ────────────────────────────────────────


@pytest.fixture
def openai_server():
    """Answers /v1/models and /v1/chat/completions like vLLM / llama-server would."""
    calls = []

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def _send(self, payload):
            body = json.dumps(payload).encode()
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            self._send({"data": [{"id": "writer"}]})

        def do_POST(self):
            request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            calls.append({"auth": self.headers.get("Authorization"), **request})
            prompt = request["messages"][-1]["content"]
            content = _fake_writer(9)(prompt)
            self._send({"choices": [{"message": {"role": "assistant", "content": content}}]})

    server = HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}", calls
    server.shutdown()


def _docs(tmp_path):
    path = tmp_path / "handbook.txt"
    path.write_text("\n\n".join(f"Rule {i}: always label box {i} clearly." for i in range(5)))
    return path


class _Upload:
    def __init__(self, path):
        self.name = str(path)


def test_ui_creates_training_data_from_a_server(openai_server, tmp_path):
    from config.constants import SYNTH_WRITERS
    from core.state import app_state
    from ui.handlers import on_synthesize

    url, calls = openai_server
    status, preview, stats, ds, jsonl = on_synthesize(
        [_Upload(_docs(tmp_path))], SYNTH_WRITERS[0], url, "", "sk-test", "", 2, 7, 10, None
    )
    assert status.startswith("✅ 2 question/answer pairs"), status
    assert len(ds) == 2 and ds.column_names == ["instruction", "output"] and len(preview) == 2
    rows = [json.loads(line) for line in open(jsonl, encoding="utf-8")]
    assert rows[0]["source"] == "handbook.txt" and rows[0]["score"] == 9
    assert calls[0]["model"] == "writer" and calls[0]["auth"] == "Bearer sk-test"
    app_state.session_for(None).release("synth")


@pytest.mark.parametrize(
    ("args", "message"),
    [((None, 0, "http://x", "", "", ""), "Upload one or more documents"),
     (("docs", 0, "", "", "", ""), "Enter the server URL"),
     (("docs", 1, "", "", "", ""), "Enter the local model")],
)  # fmt: skip
def test_ui_input_checks(tmp_path, args, message):
    from config.constants import SYNTH_WRITERS
    from ui.handlers import on_synthesize

    files = [_Upload(_docs(tmp_path))] if args[0] == "docs" else None
    status, *_ = on_synthesize(files, SYNTH_WRITERS[args[1]], *args[2:], 2, 7, 10, None)
    assert message in status


def test_cli_writes_jsonl_that_trains(openai_server, tmp_path, monkeypatch):
    """synthesize → JSONL → train: the generated file is valid training data."""
    url, _ = openai_server
    out = tmp_path / "data.jsonl"
    result = CliRunner().invoke(commands.app, ["synthesize", "--input", str(_docs(tmp_path)),
                                               "--server", url, "--output", str(out)])  # fmt: skip
    assert result.exit_code == 0, result.output
    assert len(out.read_text().splitlines()) == 2
    seen = {}
    monkeypatch.setattr(commands, "train_model",
                        lambda **kw: seen.update(rows=len(kw["dataset"])) or ("✅", []))  # fmt: skip
    result = CliRunner().invoke(commands.app, ["train", "--model", "m", "--data", str(out),
                                               "--output", str(tmp_path / "run")])  # fmt: skip
    assert result.exit_code == 0, result.output
    assert seen["rows"] == 2


@pytest.mark.parametrize(
    ("args", "message"),
    [(["--input", "a.txt"], "Give either --server"),
     (["--input", "a.txt", "--server", "http://x", "--model", "m"], "Give either --server"),
     (["--input", "missing.txt", "--server", "http://x"], "File not found")],
)  # fmt: skip
def test_cli_input_checks(args, message):
    result = CliRunner().invoke(commands.app, ["synthesize", *args])
    assert result.exit_code == 1 and message in result.stderr


def test_local_writer_runs_a_real_model(tmp_path):
    """The local path loads a real (tiny) chat model and asks it; its gibberish isn't JSON."""
    from inference.synthesize import local_writer
    from tests.test_smoke_training import TINY_CHAT_MODEL, _cached_model

    model = _cached_model(TINY_CHAT_MODEL)
    ask = local_writer(model, max_new_tokens=8)
    assert isinstance(ask("Say hi"), str)
    rows, stats = synthesize_pairs([("a.txt", "Some text.")], ask, 1, 0)
    assert rows == [] and stats["failed_chunks"] == 1


def test_documents_respect_the_ui_path_allowlist(tmp_path, monkeypatch):
    import core.state as state

    monkeypatch.setattr(state, "_allowed_roots", None)
    state.restrict_paths_to([str(tmp_path)])
    try:
        with pytest.raises(ValueError, match="Path outside"):
            read_document("/etc/hostname.txt")
    finally:
        monkeypatch.setattr(state, "_allowed_roots", None)
