"""
data/documents.py
=================
Layer 2 — read documents (PDF, Word, text, Markdown) as plain text and cut them into
overlapping chunks, the input for "Create training data" (inference/synthesize.py).
"""

import os
import re

from config.constants import (
    DOCUMENT_EXTENSIONS,
    FILE_EXT_DOCX,
    FILE_EXT_PDF,
    HAS_DOCX,
    HAS_PDF,
    SYNTH_CHUNK_CHARS,
    SYNTH_CHUNK_OVERLAP,
)
from core.state import validate_path_traversal
from data.loader import extract_text_from_pdf


def read_document(path: str) -> str:
    """Plain text of a PDF, .docx, .txt or .md file. Raises ValueError with a readable message."""
    if err := validate_path_traversal(path):
        raise ValueError(err)
    ext = os.path.splitext(path)[1].lower()
    if ext not in DOCUMENT_EXTENSIONS:
        raise ValueError(
            f"{os.path.basename(path)}: use one of {', '.join(DOCUMENT_EXTENSIONS)} documents."
        )
    if not os.path.isfile(path):
        raise ValueError(f"File not found: {path}")
    if ext == FILE_EXT_PDF:
        if not HAS_PDF:
            raise ValueError('Reading PDFs needs pypdf: pip install "pypdf>=6.16.1"')
        return extract_text_from_pdf(path)
    if ext == FILE_EXT_DOCX:
        if not HAS_DOCX:
            raise ValueError(
                'Reading Word files needs python-docx: pip install "python-docx>=1.1.0"'
            )
        import docx  # lazy

        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:  # table cells, one row per line
            parts += [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        return "\n".join(parts)
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def chunk_text(
    text: str, size: int = SYNTH_CHUNK_CHARS, overlap: int = SYNTH_CHUNK_OVERLAP
) -> list[str]:
    """Split text into chunks of about ``size`` characters, cutting at paragraph or
    sentence ends where possible; consecutive chunks share ``overlap`` characters."""
    text = re.sub(r"[ \t]+", " ", text or "")
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if not text:
        return []
    if size <= overlap:
        raise ValueError("Chunk size must be larger than the overlap.")
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):  # prefer a paragraph, then a sentence, then a word boundary
            window = text[start:end]
            for sep in ("\n\n", ". ", "\n", " "):
                cut = window.rfind(sep)
                if cut > size // 2:
                    end = start + cut + len(sep)
                    break
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        next_start = max(end - overlap, start + 1)
        space = text.find(" ", next_start, end)  # begin the overlap on a whole word
        start = space + 1 if space != -1 else next_start
    return chunks


def document_chunks(
    paths: list[str], size: int = SYNTH_CHUNK_CHARS, overlap: int = SYNTH_CHUNK_OVERLAP
) -> list[tuple[str, str]]:
    """(file name, chunk) for every chunk of every document, in order."""  # fmt: skip
    out: list[tuple[str, str]] = []
    for path in paths:
        name = os.path.basename(path)
        out += [(name, chunk) for chunk in chunk_text(read_document(path), size, overlap)]
    return out
